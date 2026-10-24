"""Converte a minuta salva em .docx formatado, com aviso de minuta no topo.

O arquivo é montado em memória a cada download a partir do texto guardado na coleção
`pecas`: nada fica gravado em disco no servidor.
"""
from __future__ import annotations

import io
import re
import unicodedata

AVISO = "MINUTA GERADA POR IA — REVISAR ANTES DE PROTOCOLAR"


class DocxIndisponivel(RuntimeError):
    """python-docx não instalado no ambiente do backend."""


def build_docx(rascunho: str) -> bytes:
    # Carregado só no download: sem o pacote, o resto do backend continua funcionando.
    try:
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError as e:
        raise DocxIndisponivel(
            "Download em .docx indisponível: instale as dependências com "
            "backend/venv/Scripts/python -m pip install -r backend/requirements.txt") from e

    d = docx.Document()
    d.styles["Normal"].font.name = "Times New Roman"
    d.styles["Normal"].font.size = Pt(12)

    aviso = d.add_paragraph()
    aviso.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = aviso.add_run(AVISO)
    run.bold = True

    for line in rascunho.splitlines():
        text = line.strip()
        if not text:
            continue
        clean = re.sub(r"^#+\s*", "", text).replace("**", "")
        is_heading = text.startswith("#") or (clean.isupper() and len(clean) < 90)
        if is_heading:
            d.add_heading(clean, level=2)
        else:
            paragrafo = d.add_paragraph(clean)
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    saida = io.BytesIO()
    d.save(saida)
    return saida.getvalue()


def nome_do_arquivo(titulo: str) -> str:
    ascii_ = unicodedata.normalize("NFKD", titulo).encode("ascii", "ignore").decode()
    base = re.sub(r"[^A-Za-z0-9]+", "-", ascii_).strip("-").lower()
    return f"minuta-{base or 'peca'}.docx"
